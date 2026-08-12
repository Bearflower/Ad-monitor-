from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from adwatch.storage.db import Database

EXPECTED_SHEETS = (
    "收支总览", "收入明细", "进货成本", "出库产品", "费用支出",
    "广告充值", "刷单成本", "月度汇总", "合伙人分润", "汇率换算",
    "前期投入",
)
INPUT_SHEETS = {
    "收入明细", "进货成本", "出库产品", "费用支出", "广告充值",
    "刷单成本", "汇率换算", "前期投入",
}


@dataclass(frozen=True)
class PreviewRow:
    sheet: str
    row_number: int
    status: str
    natural_key: str
    message: str
    normalized: dict[str, str]


@dataclass(frozen=True)
class ExcelPreviewResult:
    sheet_names: tuple[str, ...]
    rows: tuple[PreviewRow, ...]
    counts: dict[str, int]


class ExcelPrototypePreview:
    def __init__(self, database: Database) -> None:
        self.database = database

    def preview(self, workbook_path: str | Path) -> ExcelPreviewResult:
        workbook = load_workbook(
            workbook_path, data_only=False, read_only=True
        )
        names = tuple(workbook.sheetnames)
        rows: list[PreviewRow] = []
        seen: dict[str, dict[str, str]] = {}
        for sheet in workbook.worksheets:
            if sheet.title not in INPUT_SHEETS:
                continue
            headers = [self._text(cell.value) for cell in sheet[1]]
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                normalized = {
                    header: self._text(value)
                    for header, value in zip(headers, values, strict=False)
                    if header and value not in (None, "")
                    and not self._is_formula(value)
                }
                if not normalized:
                    continue
                natural_key = self._natural_key(
                    sheet.title, normalized, row_number
                )
                status, message = self._classify(
                    sheet.title, normalized, natural_key, seen
                )
                rows.append(
                    PreviewRow(
                        sheet.title, row_number, status, natural_key,
                        message, normalized,
                    )
                )
                if status in {"valid", "missing_rate"}:
                    seen[natural_key] = normalized
        counts = Counter(row.status for row in rows)
        return ExcelPreviewResult(
            names,
            tuple(rows),
            {name: counts.get(name, 0) for name in (
                "valid", "duplicate", "conflict", "missing_rate", "invalid"
            )},
        )

    def _classify(self, sheet, row, key, seen):
        if not self._date_value(row.get("日期", row.get("年月", ""))):
            return "invalid", "missing or invalid date/month"
        if key in seen:
            return (
                ("duplicate", "same natural key and values")
                if seen[key] == row
                else ("conflict", "same natural key with different values")
            )
        database_status = self._database_status(sheet, row)
        if database_status is not None:
            return database_status
        currency = row.get("币种", "")
        currency_code = {"฿": "THB", "THB": "THB", "S$": "SGD", "SGD": "SGD"}.get(currency)
        if currency_code and sheet != "汇率换算":
            month = self._month(row.get("日期", ""))
            if month and not self._has_monthly_rate(currency_code, month):
                return "missing_rate", "monthly conversion rate not stored"
        return "valid", "ready for reviewed import"

    def _database_status(self, sheet: str, row: dict[str, str]):
        if sheet != "收入明细":
            return None
        order_id = row.get("订单编号", "")
        occurred_on = row.get("日期", "")
        amount = row.get("结算金额(SGD)", "")
        if not order_id or not occurred_on or not amount:
            return None
        with self.database.connect() as connection:
            matches = connection.execute(
                """
                SELECT amount_original, currency FROM manual_income_entries
                WHERE order_id=? AND occurred_on=?
                """,
                (order_id, occurred_on),
            ).fetchall()
        if not matches:
            return None
        if any(
            item["currency"] == "SGD"
            and Decimal(item["amount_original"]) == Decimal(amount)
            for item in matches
        ):
            return "duplicate", "same record already exists in database"
        return "conflict", "natural key exists with different values"

    def _has_monthly_rate(self, currency: str, year_month: str) -> bool:
        with self.database.connect() as connection:
            return connection.execute(
                "SELECT 1 FROM monthly_exchange_rate_overrides "
                "WHERE currency=? AND year_month=?",
                (currency, year_month),
            ).fetchone() is not None

    @classmethod
    def _natural_key(cls, sheet, row, row_number):
        candidates = {
            "收入明细": ("日期", "平台", "订单编号", "结算金额(SGD)"),
            "进货成本": ("日期", "SKU", "供应商", "数量"),
            "出库产品": ("日期", "订单编号", "SKU", "数量"),
            "费用支出": ("日期", "费用类别", "说明", "原币金额"),
            "广告充值": ("日期", "币种", "充值金额", "实际到账"),
            "刷单成本": ("日期", "说明", "货款", "服务费"),
            "汇率换算": ("年月",),
            "前期投入": ("日期", "投入项目", "出资人", "金额"),
        }
        values = tuple(row.get(name, "") for name in candidates.get(sheet, ()))
        return f"{sheet}:" + "|".join(values or (str(row_number),))

    @staticmethod
    def _is_formula(value) -> bool:
        return isinstance(value, str) and value.startswith("=")

    @staticmethod
    def _text(value) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, float):
            try:
                return str(Decimal(str(value)).normalize())
            except InvalidOperation:
                pass
        return str(value).strip()

    @staticmethod
    def _month(value: str) -> str:
        return value[:7] if len(value) >= 7 else ""

    @staticmethod
    def _date_value(value: str) -> bool:
        return bool(value and (len(value) >= 7))
